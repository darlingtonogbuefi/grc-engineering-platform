"""
Excel Reader

Responsible only for reading Excel files.

Input:
    .xlsx workbook

Output:
    Raw worksheet rows

No framework-specific logic exists here.

Example:

    reader = ExcelReader()

    rows = reader.read(
        "importers/sources/CAF.xlsx",
        "CAF"
    )

    for row in rows:
        print(row)
"""

from pathlib import Path
from typing import List, Any, Optional

from openpyxl import load_workbook


class ExcelReader:
    """
    Generic Excel workbook reader.
    """

    def __init__(
        self,
        read_only: bool = True,
        data_only: bool = True
    ):
        """
        Configure workbook loading.

        read_only:
            Improves performance for large files.

        data_only:
            Returns calculated cell values instead of formulas.
        """

        self.read_only = read_only
        self.data_only = data_only


    def read(
        self,
        filename: str,
        worksheet: Optional[str] = None
    ) -> List[List[Any]]:
        """
        Read worksheet rows.

        Args:

            filename:
                Path to Excel workbook.

            worksheet:
                Worksheet name.
                If None, uses active worksheet.

        Returns:

            List of rows.

            Example:

            [
                [
                    "A1.a",
                    "A1 Governance",
                    "Board direction"
                ],
                [
                    "A1.b",
                    "A1 Governance",
                    "Roles and responsibilities"
                ]
            ]
        """

        path = Path(filename)

        if not path.exists():
            raise FileNotFoundError(
                f"Excel file not found: {filename}"
            )


        workbook = load_workbook(
            filename=path,
            read_only=self.read_only,
            data_only=self.data_only
        )


        if worksheet:

            if worksheet not in workbook.sheetnames:
                raise ValueError(
                    f"Worksheet '{worksheet}' not found. "
                    f"Available worksheets: {workbook.sheetnames}"
                )

            sheet = workbook[worksheet]

        else:

            sheet = workbook.active

            if sheet is None:
                raise ValueError(
                    "Workbook does not contain an active worksheet."
                )


        rows = []


        for row in sheet.iter_rows(
            values_only=True
        ):

            rows.append(
                list(row)
            )


        workbook.close()


        return rows
