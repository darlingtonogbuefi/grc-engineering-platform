# engine\reporting\export.py


"""
Export Engine

Exports generated reports to multiple formats.

Supported formats:

- JSON
- YAML
- CSV
- Markdown
- HTML
- Excel
- PDF

Features:

- Output directory management
- Timestamped filenames
- Report packaging
- Export manifest generation
- Structured Excel export
- Structured PDF export
"""

from __future__ import annotations

import csv
import html
import json
from dataclasses import asdict, is_dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from .report_generator import ReportDocument


# ==============================================================================
# Export Manager
# ==============================================================================


class ReportExporter:
    """
    Export report documents into multiple formats.
    """

    def __init__(
        self,
        output_directory: Path,
    ) -> None:
        self.output_directory = output_directory
        self.output_directory.mkdir(
            parents=True,
            exist_ok=True,
        )

    # ------------------------------------------------------------------
    # Export All
    # ------------------------------------------------------------------

    def export_all(
        self,
        report: ReportDocument,
        base_name: str,
    ) -> dict[str, Path]:
        """
        Export every supported format.
        """
        exports = {
            "json": self.export_json(
                report,
                base_name,
            ),
            "yaml": self.export_yaml(
                report,
                base_name,
            ),
            "csv": self.export_csv(
                report,
                base_name,
            ),
            "markdown": self.export_markdown(
                report,
                base_name,
            ),
            "html": self.export_html(
                report,
                base_name,
            ),
            "excel": self.export_excel(
                report,
                base_name,
            ),
            "pdf": self.export_pdf(
                report,
                base_name,
            ),
        }

        self.export_manifest(
            base_name,
            exports,
        )

        return exports

    # ------------------------------------------------------------------
    # JSON
    # ------------------------------------------------------------------

    def export_json(
        self,
        report: ReportDocument,
        base_name: str,
    ) -> Path:
        file = self.output_directory / f"{base_name}.json"

        with file.open(
            "w",
            encoding="utf-8",
        ) as fp:
            json.dump(
                self._serialise(report),
                fp,
                indent=2,
                ensure_ascii=False,
                default=str,
            )

        return file

    # ------------------------------------------------------------------
    # YAML
    # ------------------------------------------------------------------

    def export_yaml(
        self,
        report: ReportDocument,
        base_name: str,
    ) -> Path:
        file = self.output_directory / f"{base_name}.yml"

        with file.open(
            "w",
            encoding="utf-8",
        ) as fp:
            yaml.safe_dump(
                self._serialise(report),
                fp,
                sort_keys=False,
                allow_unicode=True,
            )

        return file

    # ------------------------------------------------------------------
    # CSV
    # ------------------------------------------------------------------

    def export_csv(
        self,
        report: ReportDocument,
        base_name: str,
    ) -> Path:
        file = self.output_directory / f"{base_name}.csv"

        with file.open(
            "w",
            newline="",
            encoding="utf-8",
        ) as fp:
            writer = csv.writer(fp)

            writer.writerow(
                [
                    "Section",
                    "Content",
                ]
            )

            for section in report.sections:
                writer.writerow(
                    [
                        section.title,
                        json.dumps(
                            section.content,
                            default=str,
                        ),
                    ]
                )

        return file

    # ------------------------------------------------------------------
    # Markdown
    # ------------------------------------------------------------------

    def export_markdown(
        self,
        report: ReportDocument,
        base_name: str,
    ) -> Path:
        file = self.output_directory / f"{base_name}.md"

        with file.open(
            "w",
            encoding="utf-8",
        ) as fp:
            fp.write(
                f"# {report.metadata.title}\n\n"
            )
            fp.write(
                f"Tenant: {report.metadata.tenant}\n\n"
            )
            fp.write(
                f"Generated: {report.metadata.generated_at}\n\n"
            )

            for section in report.sections:
                fp.write(
                    f"## {section.title}\n\n"
                )
                fp.write(
                    "```json\n"
                )
                fp.write(
                    json.dumps(
                        section.content,
                        indent=2,
                        default=str,
                    )
                )
                fp.write(
                    "\n```\n\n"
                )

        return file

    # ------------------------------------------------------------------
    # HTML
    # ------------------------------------------------------------------

    def export_html(
        self,
        report: ReportDocument,
        base_name: str,
    ) -> Path:
        file = self.output_directory / f"{base_name}.html"

        with file.open(
            "w",
            encoding="utf-8",
        ) as fp:
            fp.write(
                "<!DOCTYPE html>\n"
            )
            fp.write(
                "<html>\n<head>\n"
            )
            fp.write(
                f"<title>{html.escape(report.metadata.title)}</title>\n"
            )
            fp.write(
                '<meta charset="utf-8">\n'
            )
            fp.write(
                "</head>\n<body>\n"
            )
            fp.write(
                f"<h1>{html.escape(report.metadata.title)}</h1>\n"
            )
            fp.write(
                "<p><strong>Tenant:</strong> "
                f"{html.escape(str(report.metadata.tenant))}"
                "</p>\n"
            )
            fp.write(
                "<p><strong>Generated:</strong> "
                f"{html.escape(str(report.metadata.generated_at))}"
                "</p>\n"
            )

            for section in report.sections:
                fp.write(
                    f"<h2>{html.escape(section.title)}</h2>\n"
                )
                fp.write(
                    "<pre>"
                )
                fp.write(
                    html.escape(
                        json.dumps(
                            section.content,
                            indent=2,
                            default=str,
                        )
                    )
                )
                fp.write(
                    "</pre>\n"
                )

            fp.write(
                "</body></html>"
            )

        return file

    # ------------------------------------------------------------------
    # Excel
    # ------------------------------------------------------------------

    def export_excel(
        self,
        report: ReportDocument,
        base_name: str,
    ) -> Path:
        """
        Export report into an Excel workbook.

        Workbook contains:

        - Summary
        - One structured worksheet per report section

        Nested dictionaries and lists are rendered as tables rather
        than as raw JSON strings.
        """
        file = self.output_directory / f"{base_name}.xlsx"

        workbook = Workbook()

        summary_sheet = workbook.active

        if summary_sheet is None:
            raise RuntimeError(
                "Failed to create the Excel Summary worksheet."
            )

        summary_sheet.title = "Summary"

        # --------------------------------------------------------------
        # Workbook styling
        # --------------------------------------------------------------

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )

        section_fill = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7",
        )

        header_font = Font(
            bold=True,
            color="FFFFFF",
        )

        bold_font = Font(
            bold=True,
        )

        # --------------------------------------------------------------
        # Summary
        # --------------------------------------------------------------

        summary_rows = [
            ("Report", report.metadata.title),
            ("Tenant", report.metadata.tenant),
            (
                "Generated",
                str(report.metadata.generated_at),
            ),
            (
                "Version",
                str(report.metadata.version),
            ),
            (
                "Frameworks",
                ", ".join(
                    report.metadata.frameworks
                ),
            ),
        ]

        for row_number, (label, value) in enumerate(
            summary_rows,
            start=1,
        ):
            summary_sheet.cell(
                row=row_number,
                column=1,
                value=label,
            )

            summary_sheet.cell(
                row=row_number,
                column=2,
                value=value,
            )

            summary_sheet.cell(
                row=row_number,
                column=1,
            ).font = bold_font

        summary_sheet.column_dimensions["A"].width = 24
        summary_sheet.column_dimensions["B"].width = 80

        # --------------------------------------------------------------
        # Report sections
        # --------------------------------------------------------------

        for section in report.sections:
            sheet_name = self._excel_sheet_name(
                section.title,
                workbook.sheetnames,
            )

            sheet = workbook.create_sheet(
                title=sheet_name,
            )

            sheet["A1"] = section.title
            sheet["A1"].font = Font(
                bold=True,
                size=14,
                color="FFFFFF",
            )
            sheet["A1"].fill = header_fill

            sheet.merge_cells(
                start_row=1,
                start_column=1,
                end_row=1,
                end_column=4,
            )

            content = self._serialise(
                section.content
            )

            self._write_excel_content(
                sheet=sheet,
                content=content,
                start_row=3,
                header_fill=header_fill,
                header_font=header_font,
                section_fill=section_fill,
            )

            sheet.freeze_panes = "A3"

            for column in ("A", "B", "C", "D", "E"):
                sheet.column_dimensions[column].width = 24

        workbook.save(file)

        return file

    # ------------------------------------------------------------------
    # Excel Content Writer
    # ------------------------------------------------------------------

    @classmethod
    def _write_excel_content(
        cls,
        sheet: Any,
        content: Any,
        start_row: int,
        header_fill: PatternFill,
        header_font: Font,
        section_fill: PatternFill,
    ) -> int:
        """
        Render structured report content into Excel cells.

        Returns the next available row.
        """

        if isinstance(content, dict):
            row = start_row

            # Simple dictionary:
            #
            # key | value
            #
            # Complex dictionary:
            #
            # key | value / nested content

            simple_items: list[tuple[str, Any]] = []
            complex_items: list[tuple[str, Any]] = []

            for key, value in content.items():
                if cls._is_simple_value(value):
                    simple_items.append(
                        (
                            str(key),
                            value,
                        )
                    )
                else:
                    complex_items.append(
                        (
                            str(key),
                            value,
                        )
                    )

            if simple_items:
                sheet.cell(
                    row=row,
                    column=1,
                    value="Metric",
                )
                sheet.cell(
                    row=row,
                    column=2,
                    value="Value",
                )

                for column in (1, 2):
                    sheet.cell(
                        row=row,
                        column=column,
                    ).fill = header_fill

                    sheet.cell(
                        row=row,
                        column=column,
                    ).font = header_font

                row += 1

                for key, value in simple_items:
                    sheet.cell(
                        row=row,
                        column=1,
                        value=key,
                    )

                    sheet.cell(
                        row=row,
                        column=2,
                        value=cls._excel_value(
                            value
                        ),
                    )

                    sheet.cell(
                        row=row,
                        column=1,
                    ).font = Font(
                        bold=True
                    )

                    row += 1

                row += 1

            for key, value in complex_items:
                sheet.cell(
                    row=row,
                    column=1,
                    value=key,
                )

                sheet.cell(
                    row=row,
                    column=1,
                ).fill = section_fill

                sheet.cell(
                    row=row,
                    column=1,
                ).font = Font(
                    bold=True
                )

                row += 1

                row = cls._write_excel_content(
                    sheet=sheet,
                    content=value,
                    start_row=row,
                    header_fill=header_fill,
                    header_font=header_font,
                    section_fill=section_fill,
                )

                row += 1

            return row

        if isinstance(content, list):
            if not content:
                sheet.cell(
                    row=start_row,
                    column=1,
                    value="No records",
                )
                return start_row + 1

            # List of dictionaries -> proper table.
            if all(
                isinstance(item, dict)
                for item in content
            ):
                keys: list[str] = []

                for item in content:
                    for key in item:
                        key_string = str(key)

                        if key_string not in keys:
                            keys.append(
                                key_string
                            )

                row = start_row

                for column, key in enumerate(
                    keys,
                    start=1,
                ):
                    cell = sheet.cell(
                        row=row,
                        column=column,
                        value=key.replace(
                            "_",
                            " ",
                        ).title(),
                    )

                    cell.fill = header_fill
                    cell.font = header_font

                row += 1

                for item in content:
                    for column, key in enumerate(
                        keys,
                        start=1,
                    ):
                        value = item.get(key)

                        sheet.cell(
                            row=row,
                            column=column,
                            value=cls._excel_value(
                                value
                            ),
                        )

                    row += 1

                return row

            # List of simple values.
            sheet.cell(
                row=start_row,
                column=1,
                value="Value",
            )

            sheet.cell(
                row=start_row,
                column=1,
            ).fill = header_fill

            sheet.cell(
                row=start_row,
                column=1,
            ).font = header_font

            row = start_row + 1

            for item in content:
                sheet.cell(
                    row=row,
                    column=1,
                    value=cls._excel_value(
                        item
                    ),
                )
                row += 1

            return row

        sheet.cell(
            row=start_row,
            column=1,
            value=cls._excel_value(
                content
            ),
        )

        return start_row + 1

    # ------------------------------------------------------------------
    # PDF
    # ------------------------------------------------------------------

    def export_pdf(
        self,
        report: ReportDocument,
        base_name: str,
    ) -> Path:
        """
        Export report into a structured PDF document.

        Dictionaries are rendered as metric tables and lists of
        dictionaries are rendered as report tables.
        """
        file = self.output_directory / f"{base_name}.pdf"

        document = SimpleDocTemplate(
            str(file),
            pagesize=A4,
            rightMargin=15 * mm,
            leftMargin=15 * mm,
            topMargin=15 * mm,
            bottomMargin=15 * mm,
            title=report.metadata.title,
            author="GRC Engineering Platform",
        )

        styles = getSampleStyleSheet()

        body_style = ParagraphStyle(
            "ReportBody",
            parent=styles["BodyText"],
            fontSize=9,
            leading=12,
            alignment=TA_LEFT,
        )

        table_header_style = ParagraphStyle(
            "ReportTableHeader",
            parent=body_style,
            fontName="Helvetica-Bold",
            textColor=colors.white,
        )

        story: list[Any] = []

        # --------------------------------------------------------------
        # Title
        # --------------------------------------------------------------

        story.append(
            Paragraph(
                self._pdf_escape(
                    report.metadata.title
                ),
                styles["Title"],
            )
        )

        story.append(
            Spacer(
                1,
                12,
            )
        )

        # --------------------------------------------------------------
        # Metadata table
        # --------------------------------------------------------------

        metadata_rows = [
            [
                Paragraph(
                    "<b>Tenant</b>",
                    body_style,
                ),
                Paragraph(
                    self._pdf_escape(
                        str(report.metadata.tenant)
                    ),
                    body_style,
                ),
            ],
            [
                Paragraph(
                    "<b>Generated</b>",
                    body_style,
                ),
                Paragraph(
                    self._pdf_escape(
                        str(report.metadata.generated_at)
                    ),
                    body_style,
                ),
            ],
            [
                Paragraph(
                    "<b>Version</b>",
                    body_style,
                ),
                Paragraph(
                    self._pdf_escape(
                        str(report.metadata.version)
                    ),
                    body_style,
                ),
            ],
            [
                Paragraph(
                    "<b>Frameworks</b>",
                    body_style,
                ),
                Paragraph(
                    self._pdf_escape(
                        ", ".join(
                            report.metadata.frameworks
                        )
                    ),
                    body_style,
                ),
            ],
        ]

        metadata_table = Table(
            metadata_rows,
            colWidths=[
                35 * mm,
                135 * mm,
            ],
        )

        metadata_table.setStyle(
            TableStyle(
                [
                    (
                        "VALIGN",
                        (0, 0),
                        (-1, -1),
                        "TOP",
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.5,
                        colors.lightgrey,
                    ),
                    (
                        "BACKGROUND",
                        (0, 0),
                        (0, -1),
                        colors.HexColor(
                            "#EAF2F8"
                        ),
                    ),
                    (
                        "LEFTPADDING",
                        (0, 0),
                        (-1, -1),
                        6,
                    ),
                    (
                        "RIGHTPADDING",
                        (0, 0),
                        (-1, -1),
                        6,
                    ),
                    (
                        "TOPPADDING",
                        (0, 0),
                        (-1, -1),
                        5,
                    ),
                    (
                        "BOTTOMPADDING",
                        (0, 0),
                        (-1, -1),
                        5,
                    ),
                ]
            )
        )

        story.append(
            metadata_table
        )

        story.append(
            Spacer(
                1,
                18,
            )
        )

        # --------------------------------------------------------------
        # Sections
        # --------------------------------------------------------------

        for section in report.sections:
            story.append(
                Paragraph(
                    self._pdf_escape(
                        section.title
                    ),
                    styles["Heading2"],
                )
            )

            story.append(
                Spacer(
                    1,
                    6,
                )
            )

            content = self._serialise(
                section.content
            )

            story.extend(
                self._build_pdf_content(
                    content=content,
                    body_style=body_style,
                    table_header_style=table_header_style,
                )
            )

            story.append(
                Spacer(
                    1,
                    14,
                )
            )

        document.build(
            story
        )

        return file

    # ------------------------------------------------------------------
    # PDF Content Builder
    # ------------------------------------------------------------------

    @classmethod
    def _build_pdf_content(
        cls,
        content: Any,
        body_style: ParagraphStyle,
        table_header_style: ParagraphStyle,
    ) -> list[Any]:
        """
        Convert structured content into ReportLab flowables.
        """
        flowables: list[Any] = []

        if isinstance(content, dict):
            simple_items: list[tuple[str, Any]] = []
            complex_items: list[tuple[str, Any]] = []

            for key, value in content.items():
                if cls._is_simple_value(value):
                    simple_items.append(
                        (
                            str(key),
                            value,
                        )
                    )
                else:
                    complex_items.append(
                        (
                            str(key),
                            value,
                        )
                    )

            if simple_items:
                rows = [
                    [
                        Paragraph(
                            "Metric",
                            table_header_style,
                        ),
                        Paragraph(
                            "Value",
                            table_header_style,
                        ),
                    ]
                ]

                for key, value in simple_items:
                    rows.append(
                        [
                            Paragraph(
                                cls._pdf_escape(
                                    key.replace(
                                        "_",
                                        " ",
                                    ).title()
                                ),
                                body_style,
                            ),
                            Paragraph(
                                cls._pdf_escape(
                                    cls._display_value(
                                        value
                                    )
                                ),
                                body_style,
                            ),
                        ]
                    )

                table = Table(
                    rows,
                    colWidths=[
                        65 * mm,
                        105 * mm,
                    ],
                    repeatRows=1,
                )

                table.setStyle(
                    cls._pdf_table_style()
                )

                flowables.append(
                    table
                )

                flowables.append(
                    Spacer(
                        1,
                        8,
                    )
                )

            for key, value in complex_items:
                flowables.append(
                    Paragraph(
                        cls._pdf_escape(
                            key.replace(
                                "_",
                                " ",
                            ).title()
                        ),
                        body_style,
                    )
                )

                flowables.append(
                    Spacer(
                        1,
                        4,
                    )
                )

                flowables.extend(
                    cls._build_pdf_content(
                        content=value,
                        body_style=body_style,
                        table_header_style=table_header_style,
                    )
                )

                flowables.append(
                    Spacer(
                        1,
                        8,
                    )
                )

            return flowables

        if isinstance(content, list):
            if not content:
                flowables.append(
                    Paragraph(
                        "No records",
                        body_style,
                    )
                )
                return flowables

            if all(
                isinstance(item, dict)
                for item in content
            ):
                keys: list[str] = []

                for item in content:
                    for key in item:
                        key_string = str(key)

                        if key_string not in keys:
                            keys.append(
                                key_string
                            )

                rows = [
                    [
                        Paragraph(
                            key.replace(
                                "_",
                                " ",
                            ).title(),
                            table_header_style,
                        )
                        for key in keys
                    ]
                ]

                for item in content:
                    rows.append(
                        [
                            Paragraph(
                                cls._pdf_escape(
                                    cls._display_value(
                                        item.get(
                                            key
                                        )
                                    )
                                ),
                                body_style,
                            )
                            for key in keys
                        ]
                    )

                column_width = (
                    170 * mm
                ) / max(
                    len(keys),
                    1,
                )

                table = Table(
                    rows,
                    colWidths=[
                        column_width
                        for _ in keys
                    ],
                    repeatRows=1,
                )

                table.setStyle(
                    cls._pdf_table_style()
                )

                flowables.append(
                    table
                )

                return flowables

            rows = [
                [
                    Paragraph(
                        "Value",
                        table_header_style,
                    )
                ]
            ]

            for item in content:
                rows.append(
                    [
                        Paragraph(
                            cls._pdf_escape(
                                cls._display_value(
                                    item
                                )
                            ),
                            body_style,
                        )
                    ]
                )

            table = Table(
                rows,
                colWidths=[
                    170 * mm
                ],
                repeatRows=1,
            )

            table.setStyle(
                cls._pdf_table_style()
            )

            flowables.append(
                table
            )

            return flowables

        flowables.append(
            Paragraph(
                cls._pdf_escape(
                    cls._display_value(
                        content
                    )
                ),
                body_style,
            )
        )

        return flowables

    # ------------------------------------------------------------------
    # PDF Table Styling
    # ------------------------------------------------------------------

    @staticmethod
    def _pdf_table_style() -> TableStyle:
        """
        Consistent styling for generated PDF tables.
        """
        return TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor(
                        "#1F4E78"
                    ),
                ),
                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.white,
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.HexColor(
                        "#B7C9D6"
                    ),
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "TOP",
                ),
                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [
                        colors.white,
                        colors.HexColor(
                            "#F4F7F9"
                        ),
                    ],
                ),
            ]
        )

    # ------------------------------------------------------------------
    # Manifest
    # ------------------------------------------------------------------

    def export_manifest(
        self,
        base_name: str,
        exports: dict[str, Path],
    ) -> Path:
        manifest = {
            "generated": datetime.now(
                timezone.utc
            ).isoformat(),
            "report": base_name,
            "files": {
                key: str(value)
                for key, value in exports.items()
            },
        }

        file = (
            self.output_directory
            / f"{base_name}.manifest.json"
        )

        with file.open(
            "w",
            encoding="utf-8",
        ) as fp:
            json.dump(
                manifest,
                fp,
                indent=2,
            )

        return file

    # ------------------------------------------------------------------
    # Serialisation
    # ------------------------------------------------------------------

    @staticmethod
    def _serialise(
        obj: Any,
    ) -> Any:
        """
        Convert dataclass instances into dictionaries recursively.
        """
        if is_dataclass(obj) and not isinstance(
            obj,
            type,
        ):
            return asdict(obj)

        return obj

    # ------------------------------------------------------------------
    # Excel Helpers
    # ------------------------------------------------------------------

    @staticmethod
    def _excel_sheet_name(
        title: str,
        existing_names: list[str],
    ) -> str:
        """
        Create a valid unique Excel worksheet name.
        """
        invalid_characters = (
            "\\",
            "/",
            "*",
            "?",
            ":",
            "[",
            "]",
        )

        name = title

        for character in invalid_characters:
            name = name.replace(
                character,
                "_",
            )

        name = name[:31] or "Section"

        original = name
        counter = 1

        while name in existing_names:
            suffix = f"_{counter}"
            name = (
                original[
                    : 31 - len(suffix)
                ]
                + suffix
            )
            counter += 1

        return name

    # ------------------------------------------------------------------
    # Value Helpers
    # ------------------------------------------------------------------

    @staticmethod
    def _is_simple_value(
        value: Any,
    ) -> bool:
        """
        Determine whether a value can be represented directly in a
        metric/value table.
        """
        return (
            value is None
            or isinstance(
                value,
                (
                    str,
                    int,
                    float,
                    bool,
                ),
            )
        )

    @staticmethod
    def _excel_value(
        value: Any,
    ) -> Any:
        """
        Convert a Python value into an Excel-compatible value.
        """
        if value is None:
            return ""

        if isinstance(
            value,
            (
                str,
                int,
                float,
                bool,
            ),
        ):
            return value

        if isinstance(
            value,
            datetime,
        ):
            return value.isoformat()

        return json.dumps(
            value,
            default=str,
            ensure_ascii=False,
        )

    @staticmethod
    def _display_value(
        value: Any,
    ) -> str:
        """
        Convert a value into a readable report representation.
        """
        if value is None:
            return ""

        if isinstance(
            value,
            bool,
        ):
            return "Yes" if value else "No"

        if isinstance(
            value,
            float,
        ):
            if 0 <= value <= 1:
                return f"{value * 100:.0f}%"

            return str(value)

        return str(value)

    @staticmethod
    def _pdf_escape(
        value: str,
    ) -> str:
        """
        Escape text before inserting it into ReportLab Paragraphs.
        """
        return html.escape(
            str(value)
        )


# ==============================================================================
# Convenience API
# ==============================================================================


def export_report(
    report: ReportDocument,
    output_directory: Path,
    base_name: str | None = None,
) -> dict[str, Path]:
    """
    Export a report into all supported formats.
    """
    exporter = ReportExporter(
        output_directory
    )

    if base_name is None:
        timestamp = datetime.now(
            timezone.utc
        ).strftime(
            "%Y%m%d_%H%M%S"
        )
        base_name = (
            f"assessment_{timestamp}"
        )

    return exporter.export_all(
        report,
        base_name,
    )
